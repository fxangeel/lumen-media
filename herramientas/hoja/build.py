from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

DEEP = "082F34"
TURQ = "14B8A6"
TURQ_L = "CCFBF1"
GREY = "64748B"
INPUT = "FFF9C4"
F = "Arial"

thin = Side(style="thin", color="E2E8F0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def hdr(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=DEEP)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30

def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(name=F, bold=True, size=16, color=DEEP)
    ws["A2"] = sub
    ws["A2"].font = Font(name=F, italic=True, size=9, color=GREY)

def body(c, wrap=False, center=False, color="0F172A", bold=False):
    c.font = Font(name=F, size=10, color=color, bold=bold)
    c.alignment = Alignment(vertical="top", wrap_text=wrap, horizontal="center" if center else None)
    c.border = border

IG1 = "Hola, soy Angel 👋 y esto es Lumen Web Design.\n\nCreo páginas web modernas y minimalistas para negocios que quieren crecer: pedidos online, reservas y un diseño perfecto para el teléfono.\n\n📍 Desde Aragua, Venezuela, para todo el país y el exterior.\n\n¿Tu negocio aún no tiene web? Escríbeme por WhatsApp, el enlace está en mi biografía 💬\n\nLumen: le damos luz a tu negocio ✨"
IG2 = "Si tu negocio no se ve bien en el teléfono, ese cliente se va con la competencia en dos segundos 📉\n\nEl café del video es solo un ejemplo: tu web lleva tu logo, tus colores, tus fotos y tus productos reales 🎨\n\n¿Qué te gustaría que tuviera la web de tu negocio? Te leo en los comentarios 👇 o escríbeme por WhatsApp desde el enlace de mi biografía."
IG3 = "Todos los paquetes incluyen ✨\n📱 Diseño que se ve perfecto en celular, tablet y computadora\n💬 Botón de WhatsApp para que te escriban al instante\n🔎 Optimización básica para aparecer en Google\n🎓 Te enseño a manejar tu página\n🎁 2 semanas de mantenimiento gratis\n\n¿Cuál paquete le vendría mejor a tu negocio? Cuéntame en los comentarios 👇 o escríbeme por WhatsApp desde el enlace de mi biografía."
TK1 = "Hola, soy Angel 👋 Diseño páginas web para negocios de Venezuela y del exterior. Si tu negocio aún no tiene web, escríbeme por mensaje directo y te ayudo 💬"
TK2 = "El café del video es solo un ejemplo: tu web lleva tu logo, tus colores y tus productos reales 🎨 ¿Qué te gustaría que tuviera la web de tu negocio? Cuéntame en los comentarios o por mensaje directo 👇"
TK3 = "Todos los paquetes incluyen diseño para celular y computadora, botón de WhatsApp, optimización para Google, capacitación y 2 semanas de mantenimiento gratis ✨ ¿Cuál le vendría mejor a tu negocio? Escríbeme por mensaje directo y te hago tu cotización 💬"
IG4 = "Las redes sociales sirven para que te descubran, pero tu página web es donde te eligen 💡\n\nTu web le da a tu negocio una dirección propia que puedes poner en tu perfil, en tu tarjeta, en Google Maps y en tu estado de WhatsApp. Y trabaja las 24 horas, aunque no estés respondiendo mensajes 🕐\n\n¿Tu negocio ya tiene web o todavía vende solo por redes? Cuéntame en los comentarios 👇 o escríbeme por WhatsApp desde el enlace de mi biografía."
IG5 = "Lo que necesito para empezar tu web 📝\n📸 Tu logo y fotos de tus productos o de tu local\n🎨 Tus colores\n📋 Tu menú, servicios o lista de precios\n📍 Tu dirección, horario y número de WhatsApp\n\nSi algo te falta, no pasa nada: lo vemos juntos en la primera conversación 🤝\n\n¿Qué es lo que más te frena para tener la página web de tu negocio? Te leo en los comentarios 👇 o escríbeme por WhatsApp desde el enlace de mi biografía."
TK4 = "Tu web le da a tu negocio una dirección propia que puedes poner en tu perfil, en tu tarjeta y en Google Maps, y trabaja las 24 horas aunque no estés respondiendo mensajes 🕐 ¿Tu negocio ya tiene web o vende solo por redes? Cuéntame en los comentarios o escríbeme por mensaje directo 💬"
TK5 = "Para empezar tu web solo necesito tu logo, fotos de tus productos, tu menú o lista de precios y tu WhatsApp 📝 Si algo te falta, lo vemos juntos 🤝 ¿Qué es lo que más te frena para tener tu página web? Cuéntame en los comentarios o escríbeme por mensaje directo 💬"
IG6 = "Una web lenta no solo pierde visitas: Google también la muestra más abajo en los resultados 🔎\n\nY el botón de WhatsApp no es un detalle: es el camino más corto entre \"me interesa\" y \"quiero comprar\" 💬\n\n¿Cuál de estos 3 errores has visto más en las webs de negocios? Te leo en los comentarios 👇 o escríbeme por WhatsApp desde el enlace de mi biografía."
TK6 = "Una web lenta también aparece más abajo en Google 🔎 y el botón de WhatsApp es el camino más corto entre \"me interesa\" y \"quiero comprar\" 💬 ¿Cuál de estos 3 errores has visto más? Cuéntame en los comentarios o escríbeme por mensaje directo"
HT_IG ="#diseñoweb #paginasweb #webdesign #emprendedoresvenezolanos #aragua #maracay #venezuela #negociosvenezuela #marketingdigital #emprendimiento"

ROWS = [
    # n, video, fecha, hora, red, estado, enlace, descripcion, hashtags, archivo, portada
    (1, "Presentación de Lumen", "jue 15/10/2026", "1:50 pm", "Instagram", "Publicado",
     "https://www.instagram.com/reel/DehiSD7DYrn/", IG1, HT_IG, "lumen-reel-01-presentacion.mp4", "Logo (seg. 3)"),
    (1, "Presentación de Lumen", "jue 15/10/2026", "3:46 pm", "TikTok", "Publicado",
     "https://www.tiktok.com/@Lumen Web Design/video/7696983331849686293", TK1,
     "#diseñoweb #paginasweb #emprendedoresvenezolanos #aragua #venezuela", "lumen-reel-01-presentacion.mp4", "Primer fotograma"),
    (2, "Web moderna en el teléfono", "sáb 17/10/2026", "11:00 am", "Instagram", "Programado",
     "", IG2, HT_IG, "lumen-reel-02-web-en-el-telefono.mp4", "\"Así se ve una web moderna\""),
    (2, "Web moderna en el teléfono", "sáb 17/10/2026", "11:00 am", "TikTok", "Programado",
     "", TK2, "#diseñoweb #paginasweb #emprendedoresvenezolanos #venezuela #negocios",
     "lumen-reel-02-web-en-el-telefono.mp4", "Primer fotograma"),
    (3, "¿Cuánto cuesta una página web?", "lun 19/10/2026", "11:00 am", "Instagram", "Programado",
     "", IG3, HT_IG, "lumen-reel-03-cuanto-cuesta-una-web.mp4", "Tarjeta paquete Completo"),
    (3, "¿Cuánto cuesta una página web?", "lun 19/10/2026", "11:00 am", "TikTok", "Programado",
     "", TK3, "#diseñoweb #paginasweb #emprendedoresvenezolanos #venezuela #emprendimiento",
     "lumen-reel-03-cuanto-cuesta-una-web.mp4", "Primer fotograma"),
    (4, "¿Tu negocio solo está en redes?", "mié 21/10/2026", "11:00 am", "Instagram", "Programado",
     "", IG4, HT_IG, "lumen-reel-04-solo-en-redes.mp4", "Título (seg. 1,5)"),
    (4, "¿Tu negocio solo está en redes?", "mié 21/10/2026", "11:00 am", "TikTok", "Programado",
     "", TK4, "#diseñoweb #paginasweb #emprendedoresvenezolanos #venezuela #negocios",
     "lumen-reel-04-solo-en-redes.mp4", "Primer fotograma"),
    (5, "Así hago tu web en 4 pasos", "vie 23/10/2026", "11:00 am", "Instagram", "Programado",
     "", IG5, HT_IG, "lumen-reel-05-proceso-4-pasos.mp4", "Título (seg. 1,8)"),
    (5, "Así hago tu web en 4 pasos", "vie 23/10/2026", "11:00 am", "TikTok", "Programado",
     "", TK5, "#diseñoweb #paginasweb #emprendedoresvenezolanos #venezuela #emprendimiento",
     "lumen-reel-05-proceso-4-pasos.mp4", "Primer fotograma"),
    (6, "3 errores que espantan clientes", "vie 16/10/2026", "12:15 pm", "Instagram", "Programado",
     "", IG6, HT_IG, "lumen-reel-06-errores-que-espantan-clientes.mp4", "Título (seg. 1,8)"),
    (6, "3 errores que espantan clientes", "vie 16/10/2026", "12:15 pm", "TikTok", "Programado",
     "", TK6, "#diseñoweb #paginasweb #emprendedoresvenezolanos #venezuela #negocios",
     "lumen-reel-06-errores-que-espantan-clientes-tiktok.mp4", "Primer fotograma"),
]

wb = Workbook()

# ---------------- Calendario ----------------
ws = wb.active
ws.title = "Calendario"
title(ws, "Lumen Web Design · Calendario de publicaciones",
      "Horas en hora de Venezuela. Publicación automática con Metricool. Todos los videos van marcados como contenido hecho con IA.")
H = ["N°", "Video", "Fecha", "Hora", "Red", "Estado", "Enlace", "Descripción", "Hashtags", "Archivo (Videos\\Lumen)", "Portada"]
hdr(ws, 4, H, [5, 26, 15, 10, 11, 13, 30, 60, 34, 34, 22])
for r, row in enumerate(ROWS, start=5):
    for col, v in enumerate(row, start=1):
        c = ws.cell(row=r, column=col, value=v if v != "" else None)
        body(c, wrap=col in (2, 7, 8, 9, 10, 11), center=col in (1, 3, 4, 5, 6))
        if col == 7 and v:
            c.hyperlink = v.replace(" ", "%20")
            c.value = "Ver publicación"
            c.font = Font(name=F, size=10, color="0563C1", underline="single")
        if col == 7 and not v:
            c.value = "Se genera al publicarse"
            c.font = Font(name=F, size=9, italic=True, color=GREY)
        if col == 5:
            c.font = Font(name=F, size=10, bold=True, color=DEEP)
    ws.row_dimensions[r].height = 150 if row[4] == "Instagram" else 80
last = 4 + len(ROWS)
dv = DataValidation(type="list", formula1='"Publicado,Programado,Pendiente,Borrador"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"F5:F{last}")
ws.conditional_formatting.add(f"F5:F{last}", FormulaRule(formula=['$F5="Publicado"'],
                              fill=PatternFill("solid", fgColor=TURQ_L), font=Font(name=F, bold=True, color="0F766E")))
ws.conditional_formatting.add(f"F5:F{last}", FormulaRule(formula=['$F5="Programado"'],
                              fill=PatternFill("solid", fgColor="E0F2FE"), font=Font(name=F, bold=True, color="075985")))
ws.freeze_panes = "C5"
ws.auto_filter.ref = f"A4:K{last}"

# ---------------- Resultados ----------------
# Los números salen de analisis/estadisticas.json (Metricool), que llena herramientas/analisis/estadisticas.py.
import sys, os, json
from datetime import datetime
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "analisis"))
import estadisticas as E
EST = E.cargar(E.DATOS, E.datos_vacios())
VID = {v["n"]: v for v in E.cargar(E.VIDEOS, {"videos": []})["videos"]}
ULT = E.ultimas(EST)

def leads_de(n, red, estados=None):
    return sum(1 for l in EST["leads"] if l.get("video") == n and l["red"] == red
               and (estados is None or l.get("estado") in estados))

rs = wb.create_sheet("Resultados")
title(rs, "Resultados por publicación",
      "Claude llena esta pestaña con los datos de Metricool (a las 48 h y a los 7 días de cada video). Las columnas grises se calculan solas.")
RH = ["N°", "Video", "Red", "Estado", "Horas publicado", "Vistas", "Alcance", "Vio +3 s (IG) / completo (TT)",
      "% visto", "Tiempo medio (s)", "Me gusta", "Comentarios", "Guardados", "Compartidos", "Interacciones",
      "Tasa de interacción", "Seguidores 48 h", "Mensajes recibidos", "Clientes que cotizaron"]
hdr(rs, 4, RH, [5, 30, 11, 13, 10, 10, 10, 14, 10, 11, 10, 12, 11, 12, 13, 13, 12, 12, 13])
AUTO = "E6F7F4"
for i, row in enumerate(ROWS):
    r = 5 + i
    n, red = row[0], row[4].lower()
    vals = [f"=Calendario!A{r}", f"=Calendario!B{r}", f"=Calendario!E{r}", f"=Calendario!F{r}"]
    for col, v in enumerate(vals, start=1):
        c = rs.cell(row=r, column=col, value=v)
        body(c, center=col != 2, wrap=col == 2)
        c.font = Font(name=F, size=10, color="008000")
    l = ULT.get((n, red))
    d = l["datos"] if l else {}
    x = E.derivadas(d, VID[n]["duracion_s"]) if l else {}
    seg = E.seguidores_48h(EST, red, datetime.fromisoformat(VID[n]["publicacion"][red])) if n in VID else None
    pct3 = d.get("vio_3s_pct") if red == "instagram" else d.get("vio_completo_pct")
    datos = {5: l["horas_desde_publicacion"] if l else None, 6: d.get("vistas"), 7: d.get("alcance"),
             8: pct3 / 100 if pct3 is not None else None,
             9: x.get("pct_visto") / 100 if x.get("pct_visto") is not None else None,
             10: x.get("tiempo_medio_s"), 11: d.get("me_gusta"), 12: d.get("comentarios"),
             13: d.get("guardados"), 14: d.get("compartidos"), 17: seg,
             18: leads_de(n, red) if l or leads_de(n, red) else None,
             19: leads_de(n, red, ("cotizó", "compró")) if l or leads_de(n, red) else None}
    for col, v in datos.items():
        c = rs.cell(row=r, column=col, value=v)
        body(c, center=True)
        c.fill = PatternFill("solid", fgColor=AUTO)
        c.number_format = "0.0%" if col in (8, 9) else ("0.0" if col in (5, 10) else "#,##0")
    if red == "tiktok":
        c = rs.cell(row=r, column=13); c.value = None; c.fill = PatternFill("solid", fgColor="F1F5F9")
    c = rs.cell(row=r, column=15, value=f"=IF(COUNT(K{r}:N{r})=0,\"\",SUM(K{r}:N{r}))")
    body(c, center=True); c.fill = PatternFill("solid", fgColor="F1F5F9"); c.number_format = "#,##0"
    c = rs.cell(row=r, column=16, value=f"=IFERROR(IF(G{r}=\"\",\"\",O{r}/G{r}),\"\")")
    body(c, center=True); c.fill = PatternFill("solid", fgColor="F1F5F9"); c.number_format = "0.0%"
    rs.row_dimensions[r].height = 30
tr = 5 + len(ROWS)
rs.cell(row=tr, column=2, value="Total / promedio")
body(rs.cell(row=tr, column=2), bold=True)
for col in range(6, 20):
    L = get_column_letter(col)
    if col == 16:
        f, fmt = f"=IFERROR(IF(G{tr}=0,\"\",O{tr}/G{tr}),\"\")", "0.0%"
    elif col in (8, 9, 10):
        f, fmt = f"=IFERROR(AVERAGE({L}5:{L}{tr-1}),\"\")", ("0.0%" if col != 10 else "0.0")
    else:
        f, fmt = f"=SUM({L}5:{L}{tr-1})", "#,##0"
    c = rs.cell(row=tr, column=col, value=f)
    body(c, center=True, bold=True)
    c.number_format = fmt
    c.fill = PatternFill("solid", fgColor=TURQ_L)
note_r = tr + 2
leido = max((l["leido"] for l in EST["lecturas"]), default=None)
notes = [
    "Notas:",
    f"• Última lectura de Metricool: {leido[:16].replace('T', ' ') if leido else 'todavía sin datos (la cuenta se conectó el 15/10/2026)'}.",
    "• Vio +3 s (Instagram): % de vistas que pasaron de los 3 primeros segundos; mide el gancho. Vio completo (TikTok): % que vio todo el video.",
    "• % visto: parte del video que se ve en promedio. Interacciones = Me gusta + Comentarios + Guardados + Compartidos. Tasa = Interacciones ÷ Alcance.",
    "• Seguidores 48 h: ganados el día de la publicación y el siguiente. TikTok no informa 'Guardados' por Metricool.",
    "• Mensajes y clientes: cuéntale a Claude cuando te escriba alguien y de qué video viene; son los números que de verdad importan.",
    "• Lo aprendido de cada video está en marca/guia-videos.md.",
]
for i, n in enumerate(notes):
    c = rs.cell(row=note_r + i, column=2, value=n)
    c.font = Font(name=F, size=9, bold=(i == 0), color=GREY)
rs.freeze_panes = "C5"

# ---------------- Resumen ----------------
sm = wb.create_sheet("Resumen")
title(sm, "Resumen", "Se actualiza solo con lo que pongas en Calendario y Resultados.")
items = [
    ("Videos creados", "=MAX(Calendario!A5:A100)", "0"),
    ("Publicaciones totales", "=COUNTA(Calendario!E5:E100)", "0"),
    ("Publicadas", "=COUNTIF(Calendario!F5:F100,\"Publicado\")", "0"),
    ("Programadas", "=COUNTIF(Calendario!F5:F100,\"Programado\")", "0"),
    ("En Instagram", "=COUNTIF(Calendario!E5:E100,\"Instagram\")", "0"),
    ("En TikTok", "=COUNTIF(Calendario!E5:E100,\"TikTok\")", "0"),
    ("Vistas totales", f"=Resultados!F{tr}", "#,##0"),
    ("Seguidores ganados (48 h)", f"=Resultados!Q{tr}", "#,##0"),
    ("Mensajes recibidos", f"=Resultados!R{tr}", "#,##0"),
    ("Clientes que cotizaron", f"=Resultados!S{tr}", "#,##0"),
]
sm.column_dimensions["A"].width = 28
sm.column_dimensions["B"].width = 16
for i, (k, f, fmt) in enumerate(items):
    r = 4 + i
    a = sm.cell(row=r, column=1, value=k); body(a, bold=True, color=DEEP)
    b = sm.cell(row=r, column=2, value=f); body(b, center=True, bold=True); b.number_format = fmt
    b.fill = PatternFill("solid", fgColor=TURQ_L)
r = 4 + len(items) + 1
sm.cell(row=r, column=1, value="Créditos de voz (ElevenLabs Starter)").font = Font(name=F, bold=True, size=11, color=DEEP)
rows2 = [
    ("Créditos disponibles", 38848, "#,##0", True),
    ("Costo por video (aprox.)", 400, "#,##0", True),
    ("Videos que alcanzan", f"=IFERROR(INT(B{r+1}/B{r+2}),0)", "0", False),
]
for i, (k, v, fmt, is_input) in enumerate(rows2):
    rr = r + 1 + i
    a = sm.cell(row=rr, column=1, value=k); body(a)
    b = sm.cell(row=rr, column=2, value=v); body(b, center=True); b.number_format = fmt
    if is_input:
        b.fill = PatternFill("solid", fgColor=INPUT); b.font = Font(name=F, size=10, color="0000FF")
    else:
        b.fill = PatternFill("solid", fgColor=TURQ_L); b.font = Font(name=F, size=10, bold=True)
sm.cell(row=r + 5, column=1,
        value="Fuente: tu cuenta de ElevenLabs (15/10/2026). Voz: Diego (desde el video 4). Los créditos se renuevan cada mes.").font = Font(name=F, size=9, italic=True, color=GREY)

wb.move_sheet("Resumen", offset=-2)
wb.active = 0
wb.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "marca", "lumen-calendario-publicaciones.xlsx"))
print("ok")
